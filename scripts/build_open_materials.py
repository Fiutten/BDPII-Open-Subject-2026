from __future__ import annotations

import csv
import re
import shutil
import subprocess
from copy import deepcopy
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader, PdfWriter


ROOT = Path(__file__).resolve().parents[1]
ORIGINAL = ROOT / "material_original"
OPEN = ROOT / "material_abierto"
ASSETS = OPEN / "assets"
PRESENTATION_EDITABLES = OPEN / "02_presentaciones" / "editables"
PRESENTATION_PDFS = OPEN / "02_presentaciones" / "individuales"
PRACTICE_EDITABLES = OPEN / "03_practicas" / "editables"
PRACTICE_PDFS = OPEN / "03_practicas" / "individuales"
PROGRAMS = OPEN / "06_programas"
PACKAGE = OPEN / "99_burjc_package"

AUTHORS = [
    "Alberto Fernández Isabel",
    "Natalia Madrueño Sierro",
    "Rubén Rodríguez Fernández",
]
COURSE = "Big Data Processing II"
DEGREE = "Máster Universitario en Análisis de Datos Deportivos"
LICENSE_NAME = "Creative Commons Atribución-CompartirIgual 4.0 Internacional"
LICENSE_URL = "https://creativecommons.org/licenses/by-sa/4.0/deed.es"
BURJC_URL = "https://burjcdigital.urjc.es"


VISUAL_LABELS = {
    "batch": "Procesamiento por lotes",
    "streaming": "Procesamiento continuo de eventos",
    "hadoop": "Ecosistema Hadoop",
    "spark_batch": "Procesamiento distribuido con Spark",
    "architecture": "Arquitectura distribuida de referencia",
    "p2peda": "Arquitectura orientada a eventos y agentes",
    "queue_pubsub": "Colas y publicación-suscripción",
    "rabbitmq": "Broker RabbitMQ",
    "sqs": "Servicio de colas gestionado",
    "kafka": "Plataforma de eventos Kafka",
    "tp_te": "Tiempo de procesamiento y tiempo de evento",
    "tumbling_windows": "Ventanas de intervalo fijo",
    "hopping_windows": "Ventanas deslizantes",
    "session_windows": "Ventanas de sesión",
    "microbatch": "Microprocesamiento por lotes",
    "spark": "Motor Apache Spark",
    "flink": "Motor Apache Flink",
    "storm": "Motor Apache Storm",
    "cap_theorem": "Compromisos del teorema CAP",
    "lambda_architecture": "Arquitectura Lambda",
    "kappa_architecture": "Arquitectura Kappa",
    "consumer": "Consumidor de eventos",
    "partitions": "Particionado de un topic",
    "offset": "Gestión de offsets",
    "consumer_group": "Grupo de consumidores",
    "replication": "Replicación y tolerancia a fallos",
    "zookeeper_kraft": "Coordinación mediante KRaft",
    "table": "Modelo tabular de streaming",
    "word_count": "Ejemplo de conteo continuo",
    "streaming-arch": "Arquitectura de Structured Streaming",
    "structured-streaming": "Flujo de Structured Streaming",
    "windows": "Ventanas temporales",
    "window": "Agregación por ventana",
    "late_data": "Gestión de datos tardíos",
    "watermark_append": "Watermark en modo append",
    "watermark_update": "Watermark en modo update",
    "kibana": "Visualización operativa",
    "kafka_replication": "Replicación en Kafka",
    "kafka_consumer_group": "Coordinación de consumidores",
    "spark_skew": "Distribución de datos con skew",
    "spark_ideal": "Distribución equilibrada",
    "spark_shuffle": "Operación de shuffle",
    "kafka_grafana": "Observabilidad de Kafka",
    "spark_web_ui": "Interfaz de monitorización de Spark",
    "transforrmer": "Arquitectura conceptual de un modelo de lenguaje",
    "bbdd": "Datos deportivos y modelos generativos",
    "carrera": "Ciclo de diseño y evaluación",
    "fewshot": "Aprendizaje mediante ejemplos en contexto",
    "table_shots": "Comparación zero-shot, one-shot y few-shot",
    "cot": "Descomposición explícita de una tarea",
    "pt_lifecycle-07": "Ciclo de post-entrenamiento",
    "intro": "Arquitectura general de recuperación aumentada",
    "rag": "Flujo de recuperación y generación",
    "chunking": "Estrategias de fragmentación",
    "vectors": "Representación vectorial",
    "futbol": "Caso de recuperación aplicado al fútbol",
    "tradeoff": "Compromisos entre recuperación y generación",
    "rag2": "Pipeline RAG",
    "sevilla": "Ejemplo de consulta con contexto deportivo",
    "errors": "Fuentes de error en RAG",
    "rag3": "Evaluación del flujo RAG",
    "agents": "Arquitectura de un agente",
    "ragagents": "Integración de RAG y agentes",
    "skills": "Relación entre agentes, skills y tools",
    "agentsai": "Agente de inteligencia artificial",
    "agentsflow": "Flujo de decisión agéntico",
    "react": "Patrón ReAct",
    "orchestrator": "Orquestación de agentes",
    "states": "Estado y memoria",
    "patterns": "Patrones de arquitectura agéntica",
    "brain": "Razonamiento y planificación",
    "aibrain": "Modelo cognitivo de un agente",
    "memoryrisk": "Riesgos de memoria persistente",
    "complex": "Complejidad de un sistema agéntico",
    "loop": "Bucle de percepción y acción",
    "tools": "Herramientas externas",
    "amongus": "Coordinación multiagente",
    "cost": "Control de costes",
    "calidad": "Evaluación de calidad",
    "gym": "Caso deportivo de ejemplo",
}


def ensure_dirs() -> None:
    for directory in (
        ASSETS,
        PRESENTATION_EDITABLES,
        PRESENTATION_PDFS,
        PRACTICE_EDITABLES,
        PRACTICE_PDFS,
        PROGRAMS,
        PACKAGE,
    ):
        directory.mkdir(parents=True, exist_ok=True)


def find_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def create_cc_badge() -> Path:
    out = ASSETS / "cc_by_sa_4_0.png"
    # El pictograma procede del paquete doclicense de TeX Live, que incorpora
    # la marca oficial CC BY-SA. Se compone con la versión 4.0 para que la
    # portada no dependa de una recreación gráfica de los iconos.
    official_pdf = Path(
        "/usr/share/texlive/texmf-dist/tex/latex/doclicense/images/"
        "doclicense-CC-by-sa-88x31.pdf"
    )
    official_png = ASSETS / ".cc_by_sa_official.png"
    if official_pdf.exists():
        subprocess.run(
            ["pdftoppm", "-png", "-singlefile", "-r", "300", str(official_pdf), str(official_png.with_suffix(""))],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        official = Image.open(official_png).convert("RGBA")
    else:
        official = Image.new("RGBA", (375, 132), "white")
        fallback = ImageDraw.Draw(official)
        fallback.text((15, 38), "CC  BY  SA", fill="#111111", font=find_font(36, True))

    image = Image.new("RGBA", (1040, 300), "white")
    draw = ImageDraw.Draw(image)
    official.thumbnail((520, 210), Image.Resampling.LANCZOS)
    image.alpha_composite(official, (42, (300 - official.height) // 2))
    font_main = find_font(54, True)
    font_sub = find_font(32, False)
    draw.text((585, 76), "CC BY-SA 4.0", fill="#17365D", font=font_main)
    draw.text((585, 158), "Atribución-CompartirIgual", fill="#333333", font=font_sub)
    draw.rounded_rectangle((12, 12, 1028, 288), radius=24, outline="#17365D", width=8)
    image.save(out, dpi=(300, 300))
    official_png.unlink(missing_ok=True)
    return out


def caption_for_image(path: str) -> str:
    stem = Path(path).stem.lower()
    return VISUAL_LABELS.get(stem, stem.replace("_", " ").replace("-", " ").title())


def latex_open_preamble() -> str:
    return r"""
% --- Edición abierta URJC ---
\newcommand{\ccbadge}{\includegraphics[height=0.52cm]{images/cc_by_sa_4_0.png}}
\newcommand{\openvisual}[1]{%
\begin{tikzpicture}
  \node[draw=colorblue, rounded corners=3pt, fill=colorblue!6,
        minimum height=2.2cm, text width=0.78\linewidth, align=center,
        inner sep=8pt] {\textbf{#1}\\[0.35em]\scriptsize Esquema conceptual original de la edición abierta};
\end{tikzpicture}}
\setbeamertemplate{footline}{%
  \leavevmode%
  \hbox{%
    \begin{beamercolorbox}[wd=.82\paperwidth,ht=2.7ex,dp=1.1ex,leftskip=1em]{author in head/foot}%
      \scriptsize Material docente en abierto URJC · CC BY-SA 4.0
    \end{beamercolorbox}%
    \begin{beamercolorbox}[wd=.18\paperwidth,ht=2.7ex,dp=1.1ex,center]{date in head/foot}%
      \scriptsize\insertframenumber/\inserttotalframenumber
    \end{beamercolorbox}%
  }%
}
"""


def license_frame() -> str:
    authors = r" \\ ".join(AUTHORS)
    return rf"""

\begin{{frame}}{{Licencia y datos de publicación}}
\small
\begin{{center}}
\ccbadge
\end{{center}}
\textbf{{Material docente en abierto de la Universidad Rey Juan Carlos}}\\
\textbf{{Asignatura:}} {COURSE}\\
\textbf{{Titulación:}} {DEGREE}\\
\textbf{{Autores:}} {authors}\\
\textbf{{Fecha:}} 2026\\[0.4em]
© 2026 Equipo docente. Algunos derechos reservados. Esta obra se distribuye bajo la licencia
\href{{{LICENSE_URL}}}{{{LICENSE_NAME}}}.\\[0.4em]
Los logotipos, marcas y referencias de terceros conservan su régimen jurídico propio y quedan excluidos de la licencia, salvo indicación expresa.\\
\textbf{{Depósito institucional:}} \url{{{BURJC_URL}}}
\end{{frame}}
"""


def final_license_frame() -> str:
    return rf"""

\begin{{frame}}{{Fuentes, imágenes y reutilización}}
\small
\begin{{itemize}}
  \item El texto, los ejemplos y los esquemas de esta edición abierta se distribuyen bajo CC BY-SA 4.0.
  \item Las imágenes de procedencia no acreditada de la versión de trabajo se han retirado y sustituido por esquemas conceptuales originales.
  \item Los logotipos y marcas citados se utilizan únicamente con finalidad identificativa y quedan excluidos de la licencia.
  \item Las referencias técnicas se conservan mediante citas y enlaces a sus fuentes correspondientes.
  \item Información de licencia: \url{{{LICENSE_URL}}}
\end{{itemize}}
\end{{frame}}
"""


def transform_tex(text: str) -> str:
    text = text.replace("\\usepackage{fontawesome5}\n", "")
    # La imagen de compilación no incluye los patrones de separación de
    # palabras de babel para español.  El texto sigue codificado en UTF-8 y
    # no depende de babel, así que lo retiramos de la copia publicable para
    # que el fuente sea reproducible también en instalaciones mínimas.
    text = re.sub(r"\\usepackage(?:\[[^\]]*\])?\{babel\}\s*", "", text)
    text = text.replace("\\faPython", "Python").replace("\\faTerminal", "Terminal")

    image_pattern = re.compile(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}")
    lines = []
    for line in text.splitlines():
        if "\\titlegraphic" in line:
            lines.append(
                r"\titlegraphic{\begin{tabular}{c}\includegraphics[width=2.6cm]{images/logourjceps.png}\\[0.35em]\ccbadge\end{tabular}}"
            )
            continue
        line = image_pattern.sub(lambda match: rf"\openvisual{{{caption_for_image(match.group(1))}}}", line)
        lines.append(line)
    text = "\n".join(lines) + "\n"
    text = text.replace("\\begin{document}", latex_open_preamble() + "\n\\begin{document}", 1)

    title_pos = text.find("\\titlepage")
    if title_pos == -1:
        raise ValueError("No se encontró la portada LaTeX")
    frame_end = text.find("\\end{frame}", title_pos)
    if frame_end == -1:
        raise ValueError("No se encontró el final de la portada LaTeX")
    frame_end += len("\\end{frame}")
    text = text[:frame_end] + license_frame() + text[frame_end:]
    text = text.replace("\\end{document}", final_license_frame() + "\n\\end{document}")
    return text


def copy_presentation_source(source: Path) -> Path:
    topic = source.parent.name
    destination = PRESENTATION_EDITABLES / topic
    destination.mkdir(parents=True, exist_ok=True)
    theme_text = (source.parent / "beamercolorthemeDSLAB.sty").read_text(encoding="utf-8")
    theme_text = theme_text.replace("./images/logos/MUSA_LOGO_blanco.jpg", "./images/MUSA_LOGO_blanco.jpg")
    (destination / "beamercolorthemeDSLAB.sty").write_text(theme_text, encoding="utf-8")
    logo_source = source.parent / "images" / "logourjceps.png"
    if not logo_source.exists():
        logo_source = source.parent / "images" / "logos" / "logourjceps.png"
    (destination / "images").mkdir(exist_ok=True)
    shutil.copy2(logo_source, destination / "images" / "logourjceps.png")
    # El tema Beamer incorpora el identificador visual del máster en la
    # cabecera. Es una marca institucional (no una imagen docente reutilizada)
    # y se mantiene expresamente fuera del alcance de la licencia CC.
    musa_logo = source.parent / "images" / "MUSA_LOGO_blanco.jpg"
    if not musa_logo.exists():
        musa_logo = source.parent / "images" / "logos" / "MUSA_LOGO_blanco.jpg"
    if musa_logo.exists():
        shutil.copy2(musa_logo, destination / "images" / "MUSA_LOGO_blanco.jpg")
    shutil.copy2(ASSETS / "cc_by_sa_4_0.png", destination / "images" / "cc_by_sa_4_0.png")
    output = destination / source.name
    output.write_text(transform_tex(source.read_text(encoding="utf-8")), encoding="utf-8")
    return output


def compile_presentations() -> list[Path]:
    outputs = []
    for source in sorted((ORIGINAL / "fuentes" / "teoria").rglob("*.tex")):
        open_source = copy_presentation_source(source)
        for suffix in (".aux", ".fdb_latexmk", ".fls", ".log", ".nav", ".out", ".pdf", ".snm", ".toc", ".vrb"):
            open_source.with_suffix(suffix).unlink(missing_ok=True)
        minted = open_source.parent / f"_minted-{open_source.stem}"
        if minted.exists():
            shutil.rmtree(minted)
        subprocess.run(
            ["latexmk", "-pdf", "-shell-escape", "-interaction=nonstopmode", "-halt-on-error", open_source.name],
            cwd=open_source.parent,
            check=True,
        )
        compiled = open_source.with_suffix(".pdf")
        target = PRESENTATION_PDFS / compiled.name
        shutil.copy2(compiled, target)
        outputs.append(target)
        subprocess.run(["latexmk", "-c", open_source.name], cwd=open_source.parent, check=False)
        for suffix in (".aux", ".fdb_latexmk", ".fls", ".log", ".nav", ".out", ".pdf", ".pyg", ".snm", ".toc", ".vrb"):
            open_source.with_suffix(suffix).unlink(missing_ok=True)
        minted = open_source.parent / f"_minted-{open_source.stem}"
        if minted.exists():
            shutil.rmtree(minted)
    return outputs


def remove_docx_image_by_target(doc: Document, target_filename: str) -> None:
    relationship_ids = {
        rel_id
        for rel_id, rel in doc.part.rels.items()
        if str(getattr(rel, "target_ref", "")).endswith(target_filename)
    }
    for blip in list(doc._element.xpath(".//a:blip")):
        rel_id = blip.get(qn("r:embed"))
        if rel_id not in relationship_ids:
            continue
        drawing = blip
        while drawing is not None and drawing.tag != qn("w:drawing"):
            drawing = drawing.getparent()
        if drawing is not None and drawing.getparent() is not None:
            drawing.getparent().remove(drawing)
    # Eliminar también la relación para que python-docx no conserve la imagen
    # como recurso huérfano dentro del DOCX publicado.
    for rel_id in relationship_ids:
        doc.part.drop_rel(rel_id)


def set_run_font(run, size: float, bold: bool = False, color: str = "222222") -> None:
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_centered(doc: Document, text: str, size: float, bold: bool = False, color: str = "222222"):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(paragraph.add_run(text), size, bold, color)
    return paragraph


def add_frontmatter(doc: Document, title: str, badge: Path) -> None:
    body = doc._body._element
    original_first = body[0] if len(body) else None
    created = []

    start = len(body)
    add_centered(doc, "MATERIAL DOCENTE EN ABIERTO", 13, True, "17365D")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(32)
    add_centered(doc, title, 24, True, "17365D")
    add_centered(doc, COURSE, 16, True, "1F4E79")
    add_centered(doc, DEGREE, 12, False)
    add_centered(doc, "Curso 2026-2027", 11, False, "666666")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(18)
    add_centered(doc, "Alberto Fernández Isabel", 10.5)
    add_centered(doc, "Natalia Madrueño Sierro", 10.5)
    add_centered(doc, "Rubén Rodríguez Fernández", 10.5)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(badge), width=Inches(2.65))
    add_centered(doc, "Depósito institucional  BURJC Digital", 10, False, "1F4E79")
    add_centered(doc, BURJC_URL, 9, False, "1F4E79")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    heading = doc.add_paragraph()
    set_run_font(heading.add_run("Información de licencia y reutilización"), 18, True, "17365D")
    p = doc.add_paragraph("Material docente en abierto de la Universidad Rey Juan Carlos.")
    p.paragraph_format.space_after = Pt(10)
    p = doc.add_paragraph(
        "© 2026 Alberto Fernández Isabel, Natalia Madrueño Sierro y Rubén Rodríguez Fernández. "
        "Algunos derechos reservados. Esta obra se distribuye bajo la licencia Creative Commons "
        "Atribución-CompartirIgual 4.0 Internacional, disponible en " + LICENSE_URL + "."
    )
    p.paragraph_format.space_after = Pt(10)
    doc.add_paragraph(
        "La licencia se aplica al contenido original de la obra. Los logotipos, marcas, citas y materiales de "
        "terceros conservan su régimen jurídico propio y quedan excluidos de la licencia, salvo indicación expresa."
    )
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    created.extend(list(body)[start:])
    for element in created:
        body.remove(element)
    insertion_index = 0 if original_first is not None else len(body)
    for offset, element in enumerate(created):
        body.insert(insertion_index + offset, element)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "Material docente en abierto URJC · CC BY-SA 4.0"
        for run in footer.runs:
            set_run_font(run, 8, False, "666666")


def create_open_practice_docx(source: Path, badge: Path) -> Path:
    doc = Document(source)
    if source.name == "BDPII_Proyecto_Final.docx":
        remove_docx_image_by_target(doc, "image1.png")
    title_map = {
        "BDPII_Tema_02_Enunciado_Apache_Kafka.docx": "Ejercicios de Apache Kafka",
        "BDPII_Tema_02_Enunciado_Spark_Structured_Streaming.docx": "Ejercicios de Spark Structured Streaming",
        "BDPII_Tema_04_Ejercicios_Sesion_01.docx": "Ejercicios de Prompt Engineering y Post-training",
        "BDPII_Tema_04_Ejercicios_Sesion_02.docx": "Ejercicios de Retrieval-Augmented Generation",
        "BDPII_Tema_04_Ejercicios_Sesion_03.docx": "Ejercicios de sistemas basados en agentes",
        "BDPII_Proyecto_Final.docx": "Proyecto final transversal",
    }
    add_frontmatter(doc, title_map[source.name], badge)
    target = PRACTICE_EDITABLES / source.name
    doc.save(target)
    return target


def prepare_practice_editables(badge: Path) -> list[Path]:
    return [
        create_open_practice_docx(source, badge)
        for source in sorted((ORIGINAL / "fuentes" / "practica").rglob("*.docx"))
    ]


def merge_pdfs(paths: list[Path], output: Path) -> None:
    writer = PdfWriter()
    for path in paths:
        reader = PdfReader(path)
        writer.append(reader, outline_item=path.stem.replace("_", " "))
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("wb") as stream:
        writer.write(stream)


def prepare_programs() -> None:
    code_root = PROGRAMS / "codigo"
    if code_root.exists():
        shutil.rmtree(code_root)
    code_root.mkdir(parents=True)
    for topic in sorted((ORIGINAL / "programas").iterdir()):
        if topic.is_dir():
            shutil.copytree(topic, code_root / topic.name, ignore=shutil.ignore_patterns("__pycache__", "*.pyc", ".env"))

    license_text = """MIT License

Copyright (c) 2026 Alberto Fernández Isabel, Natalia Madrueño Sierro y Rubén Rodríguez Fernández

Permission is hereby granted, free of charge, to any person obtaining a copy
of this software and associated documentation files (the \"Software\"), to deal
in the Software without restriction, including without limitation the rights
to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
copies of the Software, and to permit persons to whom the Software is
furnished to do so, subject to the following conditions:

The above copyright notice and this permission notice shall be included in all
copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED \"AS IS\", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
SOFTWARE.
"""
    (PROGRAMS / "LICENSE").write_text(license_text, encoding="utf-8")
    readme = f"""# Programas de ordenador de Big Data Processing II

Material docente en abierto de la Universidad Rey Juan Carlos.

- Asignatura: {COURSE}
- Titulación: {DEGREE}
- Curso: 2026-2027
- Autores: {', '.join(AUTHORS)}
- Repositorio público del código: https://github.com/Fiutten/BDPII-Open-Subject-Code-URJC
- Repositorio maestro de materiales: https://github.com/Fiutten/BDPII-Open-Subject-2026
- Depósito institucional: {BURJC_URL}
- Licencia del código: MIT

## Contenido

El directorio `codigo/` reúne programas y tres notebooks de los temas 1, 2, 4 y 5. Incluye ejemplos de procesamiento directo y streaming, productor Kafka, entorno Spark, ejercicios RAG y de agentes y una práctica Text-to-SQL con Spark y LangGraph.

Cada bloque conserva sus dependencias e instrucciones. Los archivos `.env` con credenciales no se distribuyen; solo se incluye `.env.example` cuando procede.

## Software Heritage

La versión candidata se archivará en Software Heritage cuando se cierre el conjunto de materiales. El SWHID verificado se incorporará a este documento y a la descripción del depósito antes de la entrega a BURJC Digital.

## Uso de herramientas generativas

Se utilizaron herramientas generativas de texto como apoyo editorial para organizar el repositorio, normalizar documentación y comprobar coherencia. No sustituyeron el diseño docente ni la revisión técnica del código. Todo el contenido publicado ha sido revisado y validado por el equipo autor, que asume plena responsabilidad sobre la versión final.
"""
    (PROGRAMS / "README.md").write_text(readme, encoding="utf-8")


def write_rights_matrix() -> None:
    rows = []
    image_root = ORIGINAL / "fuentes" / "teoria"
    for path in sorted(image_root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
            continue
        relative = path.relative_to(ROOT).as_posix()
        name = path.name.lower()
        if name == "logourjceps.png" or "musa" in name or name in {"urjc.png", "origin.jpg"}:
            classification = "Logotipo o marca institucional"
            action = "Conservar solo cuando identifica la institución; excluir expresamente de CC BY-SA"
            status = "Excluido de la licencia"
        else:
            classification = "Imagen docente sin licencia reutilizable verificada"
            action = "Retirar de la edición abierta y sustituir por esquema original"
            status = "No redistribuida"
        rows.append((relative, classification, status, action))
    output = OPEN / "02_presentaciones" / "MATRIZ_DERECHOS_IMAGENES.csv"
    with output.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(("archivo_original", "clasificacion", "estado_edicion_abierta", "accion"))
        writer.writerows(rows)


def main() -> None:
    ensure_dirs()
    badge = create_cc_badge()
    presentations = compile_presentations()
    merge_pdfs(presentations, OPEN / "02_presentaciones" / "Presentaciones_Big_Data_Processing_II.pdf")
    practices = prepare_practice_editables(badge)
    prepare_programs()
    write_rights_matrix()
    print(f"CC badge: {badge}")
    print(f"Presentations: {len(presentations)}")
    print(f"Practice editables: {len(practices)}")


if __name__ == "__main__":
    main()
