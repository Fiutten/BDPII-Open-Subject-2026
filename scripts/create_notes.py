#!/usr/bin/env python3
"""Genera los apuntes editables de Big Data Processing II desde Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "material_abierto" / "01_apuntes" / "Apuntes_Big_Data_Processing_II.md"
OUT = ROOT / "material_abierto" / "01_apuntes" / "editables" / "Apuntes_Big_Data_Processing_II.docx"
BADGE = ROOT / "material_abierto" / "assets" / "cc_by_sa_4_0.png"

AUTHORS = "Alberto Fernández Isabel · Natalia Madrueño Sierro · Rubén Rodríguez Fernández"
DEGREE = "Máster Universitario en Análisis de Datos Deportivos"
BLUE = "17365D"
MID_BLUE = "1F4E79"
LIGHT_BLUE = "DCE6F1"
GRAY = "666666"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    rel = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), MID_BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend((color, underline))
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text: str) -> None:
    # Markdown inline mínimo: negrita, cursiva y enlaces visibles.
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|https?://\S+)")
    for token in filter(None, pattern.split(text)):
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("*") and token.endswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith("http://") or token.startswith("https://"):
            add_hyperlink(paragraph, token.rstrip(".,"), token.rstrip(".,"))
            if token[-1:] in ".,":
                paragraph.add_run(token[-1])
        else:
            paragraph.add_run(token)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13

    for name, size, color, before, after in (
        ("Title", 25, "000000", 0, 14),
        ("Heading 1", 18, "000000", 16, 8),
        ("Heading 2", 14, "000000", 12, 5),
        ("Heading 3", 11.5, "000000", 9, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Code" not in doc.styles:
        code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = doc.styles["Code"]
    code.font.name = "DejaVu Sans Mono"
    code.font.size = Pt(8.2)
    code.font.color.rgb = RGBColor(32, 32, 32)
    code.paragraph_format.left_indent = Cm(0.55)
    code.paragraph_format.right_indent = Cm(0.35)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(7)
    code.paragraph_format.keep_together = True
    ppr = code.element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F6F9")
    ppr.append(shd)


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.1)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.2)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

        header = section.header.paragraphs[0]
        header.text = "Big Data Processing II · Apuntes"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(GRAY)

        footer = section.footer.paragraphs[0]
        footer.text = "Material docente en abierto URJC · CC BY-SA 4.0 · "
        for run in footer.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(GRAY)
        add_page_number(footer)


def cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("MATERIAL DOCENTE EN ABIERTO\nDE LA UNIVERSIDAD REY JUAN CARLOS")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(MID_BLUE)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Apuntes de Big Data Processing II")

    for text, size, bold, color in (
        (DEGREE, 14, True, BLUE),
        ("Curso 2026-2027", 11, False, GRAY),
        (AUTHORS, 10.5, False, "222222"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10 if size != 10.5 else 18)
        r = p.add_run(text)
        r.font.name = "Aptos"
        r.font.size = Pt(size)
        r.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)

    if BADGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(BADGE), width=Inches(2.75))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.add_run("Depósito institucional BURJC Digital\n").bold = True
    add_hyperlink(p, "https://burjcdigital.urjc.es", "https://burjcdigital.urjc.es")
    doc.add_page_break()


def copyright_page(doc: Document) -> None:
    doc.add_heading("Información de publicación", level=1)
    doc.add_paragraph("Material docente en abierto de la Universidad Rey Juan Carlos.")
    doc.add_paragraph(
        "© 2026 Alberto Fernández Isabel, Natalia Madrueño Sierro y Rubén Rodríguez Fernández. "
        "Algunos derechos reservados."
    )
    p = doc.add_paragraph("Este documento se distribuye bajo la licencia ")
    add_hyperlink(
        p,
        "Creative Commons Atribución-CompartirIgual 4.0 Internacional",
        "https://creativecommons.org/licenses/by-sa/4.0/deed.es",
    )
    doc.add_paragraph(
        "La licencia se aplica al texto, los ejemplos, las tablas y los esquemas originales. "
        "Los nombres, logotipos, marcas, citas y materiales de terceros conservan su régimen jurídico "
        "propio y quedan excluidos de la licencia, salvo indicación expresa."
    )
    doc.add_paragraph(
        "Esta edición no incorpora fotografías ni ilustraciones de terceros. Los fragmentos de código "
        "se presentan con finalidad docente y forman parte del contenido original de los apuntes."
    )
    doc.add_paragraph(
        "Asignatura: Big Data Processing II. Titulación: Máster Universitario en Análisis de Datos "
        "Deportivos. Lugar previsto de depósito: BURJC Digital. Fecha de esta edición: 2026."
    )
    doc.add_page_break()


def toc(doc: Document, headings: list[tuple[int, str]]) -> None:
    doc.add_heading("Índice", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(2.0)
    table.columns[1].width = Cm(14.1)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Nivel"
    hdr[1].text = "Contenido"
    for cell in hdr:
        shade(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for level, text in headings:
        if level != 1:
            continue
        cells = table.add_row().cells
        cells[0].text = "Tema" if re.match(r"^\d+\.", text) else "Sección"
        cells[1].text = text
        if len(table.rows) % 2 == 0:
            for cell in cells:
                shade(cell, "F3F6F9")
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    doc.add_page_break()


def parse_source() -> tuple[list[str], list[tuple[int, str]]]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    body_start = next(i for i, line in enumerate(lines) if line.strip() == "## Presentación")
    body = lines[body_start:]
    headings = []
    for line in body:
        match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if match:
            headings.append((len(match.group(1)), match.group(2)))
    return body, headings


def body(doc: Document, lines: list[str]) -> None:
    in_code = False
    code_lines: list[str] = []
    first_h1 = True
    for raw in lines:
        line = raw.rstrip()
        if line.startswith("~~~"):
            if in_code:
                p = doc.add_paragraph(style="Code")
                p.add_run("\n".join(code_lines))
                code_lines.clear()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not line or line == "---":
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1:
                if not first_h1:
                    doc.add_page_break()
                first_h1 = False
            doc.add_heading(title, level=level)
            continue
        bullet = re.match(r"^-\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, bullet.group(1))
            continue
        number = re.match(r"^\d+\.\s+(.+)$", line)
        if number:
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, number.group(1))
            continue
        p = doc.add_paragraph()
        add_rich_text(p, line)


def main() -> None:
    lines, headings = parse_source()
    doc = Document()
    configure_styles(doc)
    configure_sections(doc)
    cover(doc)
    copyright_page(doc)
    toc(doc, headings)
    body(doc, lines)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
